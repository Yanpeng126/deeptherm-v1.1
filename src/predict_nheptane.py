from __future__ import annotations

import argparse
import math
import re
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import lightning.pytorch as pl
import numpy as np
import torch
from chemprop.data import MoleculeDatapoint, MoleculeDataset, build_dataloader
from chemprop.featurizers import SimpleMoleculeMolGraphFeaturizer
from openpyxl import Workbook
from openpyxl.styles import Font
from rdkit import Chem
from rdkit.Chem import rdMolDescriptors

from dataset import TARGET_COLS, morgan_fp
from model import load_deeptherm


R_CAL = 1.98720425864083
TEMPERATURES = (300.0, 400.0, 500.0, 600.0, 800.0, 1000.0, 1500.0)
FLOAT_RE = re.compile(r"[+-]\d\.\d+E[+-]\d+")

ENSEMBLE_RUNS = [
    ("exp04_legacy600_w5_e150", 1, 0.0),
    ("exp06_legacy600_w20_e150", 1, 0.0),
    ("exp09_legacy600_w20_e150_seed44", 1, 0.0),
    ("exp_w20_e150_seed_46", 1, 0.0),
    ("exp_w20_e150_seed_47", 1, 0.0),
    ("exp_w20_e150_seed_48", 1, 0.0),
    ("exp_w20_e150_seed_49", 1, 0.0),
    ("exp_w20_e150_seed_50", 1, 0.0),
    ("transfer_qm9_seed42_w20_f2_e150", 2, 0.2),
    ("transfer_qm9_seed42_w20_f3_e120", 3, 0.2),
]


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data4", required=True, type=Path)
    parser.add_argument("--data5", required=True, type=Path)
    parser.add_argument("--runs-root", type=Path,
                        default=Path("runs/hf2118_mixed_ensemble"))
    parser.add_argument("--out", type=Path,
                        default=Path("results/nheptane_comparison.xlsx"))
    parser.add_argument("--batch-size", type=int, default=64)
    parser.add_argument("--num-workers", type=int, default=0)
    return parser.parse_args()


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def column_index(reference: str) -> int:
    letters = re.match(r"[A-Z]+", reference).group(0)
    result = 0
    for char in letters:
        result = result * 26 + ord(char) - ord("A") + 1
    return result - 1


def read_data4(path: Path):
    with ZipFile(path) as archive:
        shared_strings = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in root.iter():
                if local_name(item.tag) == "si":
                    shared_strings.append("".join(
                        node.text or "" for node in item.iter()
                        if local_name(node.tag) == "t"
                    ))

        sheet_name = next(
            name for name in archive.namelist()
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        )
        root = ElementTree.fromstring(archive.read(sheet_name))
        rows = []
        for row_node in root.iter():
            if local_name(row_node.tag) != "row":
                continue
            values = [None] * 10
            for cell in row_node:
                if local_name(cell.tag) != "c":
                    continue
                index = column_index(cell.attrib["r"])
                if index >= len(values):
                    continue
                cell_type = cell.attrib.get("t")
                if cell_type == "inlineStr":
                    value = "".join(
                        node.text or "" for node in cell.iter()
                        if local_name(node.tag) == "t"
                    )
                else:
                    value_node = next(
                        (node for node in cell if local_name(node.tag) == "v"),
                        None,
                    )
                    if value_node is None:
                        continue
                    value = value_node.text
                    if cell_type == "s":
                        value = shared_strings[int(value)]
                    elif index > 0:
                        value = float(value)
                values[index] = value
            rows.append(values)

    if not rows or rows[0][0] != "SMILES":
        raise ValueError("Data 4 does not contain the expected SMILES table")
    return rows[1:]


def parse_formula(value: str):
    counts = []
    for element, count in re.findall(r"([A-Z][a-z]?)\s*(\d*)", value):
        counts.append((element, int(count) if count else 1))
    return tuple(sorted(counts))


def smiles_formula(smiles: str):
    molecule = Chem.MolFromSmiles(smiles)
    if molecule is None:
        raise ValueError(f"could not parse Data 4 SMILES: {smiles}")
    return parse_formula(rdMolDescriptors.CalcMolFormula(molecule))


def nasa_values(coefficients):
    a1, a2, a3, a4, a5, a6, a7 = coefficients
    temperature = 298.15
    h_over_rt = (
        a1 + a2 * temperature / 2 + a3 * temperature**2 / 3
        + a4 * temperature**3 / 4 + a5 * temperature**4 / 5
        + a6 / temperature
    )
    s_over_r = (
        a1 * math.log(temperature) + a2 * temperature
        + a3 * temperature**2 / 2 + a4 * temperature**3 / 3
        + a5 * temperature**4 / 4 + a7
    )
    values = [
        h_over_rt * R_CAL * temperature / 1000,
        s_over_r * R_CAL,
    ]
    for temperature in TEMPERATURES:
        cp_over_r = (
            a1 + a2 * temperature + a3 * temperature**2
            + a4 * temperature**3 + a5 * temperature**4
        )
        values.append(cp_over_r * R_CAL)
    return np.asarray(values, dtype=np.float64)


def read_data5(path: Path):
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    if len(lines) % 4:
        raise ValueError("Data 5 line count is not a multiple of four")

    species = []
    seen = set()
    for offset in range(0, len(lines), 4):
        header, line2, line3, line4 = lines[offset:offset + 4]
        name = header[:24].strip()
        if name in seen:
            continue
        coefficients = [
            float(value) for value in FLOAT_RE.findall(line2 + line3 + line4)
        ]
        if len(coefficients) != 14:
            raise ValueError(f"invalid NASA record at line {offset + 1}")
        species.append({
            "name": name,
            "formula": parse_formula(header[24:44]),
            "values": nasa_values(np.asarray(coefficients[7:])),
        })
        seen.add(name)
    return species


def match_species(data4_rows, data5_species):
    source_values = np.asarray([row[1:] for row in data4_rows], dtype=np.float64)
    indices_by_formula = {}
    for index, row in enumerate(data4_rows):
        formula = smiles_formula(str(row[0]))
        indices_by_formula.setdefault(formula, []).append(index)

    scale = np.asarray([0.01, 0.01, 0.5, 0.5, 0.5, 0.5, 0.5, 0.5, 0.5])
    matched = []
    for entry in data5_species:
        candidates = np.asarray(indices_by_formula.get(entry["formula"], []),
                                dtype=int)
        if len(candidates) == 0:
            raise ValueError(
                f"no Data 4 candidate for {entry['name']} ({entry['formula']})"
            )
        differences = (source_values[candidates] - entry["values"]) / scale
        scores = np.sqrt(np.mean(differences**2, axis=1))
        best = int(candidates[np.argmin(scores)])
        hf_error = abs(source_values[best, 0] - entry["values"][0])
        s_error = abs(source_values[best, 1] - entry["values"][1])
        if hf_error > 0.02 or s_error > 0.02:
            raise ValueError(f"uncertain Data 4 match for {entry['name']}")
        matched.append({
            "species": entry["name"],
            "smiles": str(data4_rows[best][0]),
            "original": source_values[best],
        })
    return matched


def make_prediction_dataset(smiles):
    points = [
        MoleculeDatapoint.from_smi(
            value,
            y=np.zeros(len(TARGET_COLS), dtype=np.float32),
            x_d=morgan_fp(value, radius=2, n_bits=1024),
        )
        for value in smiles
    ]
    return MoleculeDataset(
        points,
        featurizer=SimpleMoleculeMolGraphFeaturizer(),
    )


def checkpoint_path(run_dir: Path):
    candidates = list(run_dir.glob("lightning_logs/version_*/checkpoints/best.ckpt"))
    if not candidates:
        raise FileNotFoundError(f"no best.ckpt found in {run_dir}")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def predict_ensemble(matched, runs_root: Path, batch_size: int,
                     num_workers: int):
    unique_smiles = list(dict.fromkeys(row["smiles"] for row in matched))
    dataset = make_prediction_dataset(unique_smiles)
    loader = build_dataloader(
        dataset,
        batch_size=batch_size,
        num_workers=num_workers,
        shuffle=False,
    )
    trainer = pl.Trainer(
        accelerator="auto",
        devices=1,
        deterministic=True,
        logger=False,
        enable_checkpointing=False,
        enable_progress_bar=False,
    )

    predictions = []
    validation_maes = []
    run_names = []
    for run_name, ffn_layers, dropout in ENSEMBLE_RUNS:
        run_dir = runs_root / run_name
        run_data = np.load(run_dir / "predictions.npz", allow_pickle=True)
        model = load_deeptherm(
            checkpoint_path(run_dir),
            ffn_layers=ffn_layers,
            dropout=dropout,
        )
        batches = trainer.predict(model, loader)
        prediction = torch.cat(batches).cpu().numpy()
        slope, intercept = run_data["hf_calibration"]
        prediction[:, 0] = slope * prediction[:, 0] + intercept
        predictions.append(prediction)
        validation_maes.append(float(np.abs(
            run_data["val_preds"] - run_data["val_truths"]
        ).mean()))
        run_names.append(run_name)

    validation_maes = np.asarray(validation_maes)
    inverse = 1.0 / validation_maes
    weights = inverse / inverse.sum()
    stacked = np.stack(predictions)
    ensemble = (stacked * weights[:, None, None]).sum(axis=0)
    prediction_by_smiles = dict(zip(unique_smiles, ensemble))
    for row in matched:
        row["prediction"] = prediction_by_smiles[row["smiles"]]
    return run_names, validation_maes, weights


def write_workbook(path: Path, matched, run_names, validation_maes, weights):
    original = np.stack([row["original"] for row in matched])
    prediction = np.stack([row["prediction"] for row in matched])
    errors = prediction - original
    mae = np.abs(errors).mean(axis=0)
    rmse = np.sqrt((errors**2).mean(axis=0))

    workbook = Workbook()
    statistics = workbook.active
    statistics.title = "Statistics"
    statistics.append(["Target", "Units", "Species count", "MAE", "RMSE"])
    for index, target in enumerate(TARGET_COLS):
        units = "kcal/mol" if index == 0 else "cal/(mol K)"
        statistics.append([
            target, units, len(matched), float(mae[index]), float(rmse[index]),
        ])

    comparison = workbook.create_sheet("Comparison")
    headers = ["Mechanism species", "SMILES"]
    for target in TARGET_COLS:
        headers.extend([f"Data 4 {target}", f"Reproduced {target}"])
    comparison.append(headers)
    for row in matched:
        values = [row["species"], row["smiles"]]
        for source, predicted in zip(row["original"], row["prediction"]):
            values.extend([float(source), float(predicted)])
        comparison.append(values)

    ensemble = workbook.create_sheet("Ensemble")
    ensemble.append(["Run", "Validation MAE", "Ensemble weight"])
    for name, val_mae, weight in zip(run_names, validation_maes, weights):
        ensemble.append([name, float(val_mae), float(weight)])

    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    for row in statistics.iter_rows(min_row=2, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = "0.000"
    for row in comparison.iter_rows(min_row=2, min_col=3):
        for cell in row:
            cell.number_format = "0.000"
    for row in ensemble.iter_rows(min_row=2, min_col=2, max_col=3):
        for cell in row:
            cell.number_format = "0.0000"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main():
    args = parse_args()
    torch.set_float32_matmul_precision("medium")
    data4_rows = read_data4(args.data4)
    data5_species = read_data5(args.data5)
    matched = match_species(data4_rows, data5_species)
    run_names, validation_maes, weights = predict_ensemble(
        matched,
        args.runs_root,
        args.batch_size,
        args.num_workers,
    )
    write_workbook(args.out, matched, run_names, validation_maes, weights)
    print(f"species={len(matched)}")
    print(f"saved to {args.out}")


if __name__ == "__main__":
    main()
